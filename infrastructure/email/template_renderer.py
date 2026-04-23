#atrpt/infrastructure/email/template_renderer.py

"""gera HTML a partir de um .docx e substitui placeholders.
    Não suporta:
        listas Word, tabelas, imagens, links, estilos complexos, múltiplos níveis de formatação
"""
from docx import Document
import re


def render_template(template_path, row=None):

    doc = Document(template_path)

    html_parts = []
    subject = None

    for p in doc.paragraphs:

        full_text = "".join(run.text for run in p.runs).strip()

        # capturar assunto
        if subject is None and full_text.lower().startswith("assunto"):
            subject = full_text.split(":", 1)[1].strip()
            continue

        indent = p.paragraph_format.left_indent
        indent_html = ""

        if indent:
            indent_html = "&nbsp;" * int(indent.pt / 2)

        line = indent_html

        for run in p.runs:

            text = run.text

            if row is not None:

                def repl(match):
                    campo = match.group(1).lower()
                    return str(row.get(campo, ""))

                text = re.sub(r"«(.*?)»", repl, text)

            if run.bold:
                text = f"<b>{text}</b>"

            line += text

        html_parts.append(line)

    html = "<br>".join(html_parts)

    return subject, html
